"""Excel workbook writer (one Components tab + one Entities tab per app, plus a Summary tab)."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import AppDocumentation

COMPONENT_HEADERS = ["Component Name", "Component Type", "Logical / Unique Name", "Parent Entity", "Component Id", "Discovery Source", "Details"]
ENTITY_HEADERS = ["Logical Name", "Display Name", "Discovery Source", "IsCustomEntity", "All Discovery Sources", "Is Activity", "Schema Name", "Object Type Code", "Depth", "Referenced By"]

_INVALID = re.compile(r"[\[\]:*?/\\]")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def sheet_name(base: str, suffix: str, used: set[str]) -> str:
    clean = _INVALID.sub(" ", base).strip()
    limit = 31 - len(suffix) - 3
    name = f"{clean[:limit].rstrip()} - {suffix}"
    candidate, n = name, 2
    while candidate.lower() in used:
        tail = f" ({n})"
        candidate = f"{name[:31 - len(tail)]}{tail}"
        n += 1
    used.add(candidate.lower())
    return candidate


def _table_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[^A-Za-z0-9_]", "_", name) or "Table"
    if base[0].isdigit():
        base = "T_" + base
    candidate, n = base, 2
    while candidate in used:
        candidate = f"{base}_{n}"
        n += 1
    used.add(candidate)
    return candidate


def _write_table(ws, headers: list[str], rows: list[list], table_name: str) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(headers))
    last_row = max(len(rows) + 1, 2)
    if rows:
        table = Table(displayName=table_name, ref=f"A1:{last_col}{last_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for idx, header in enumerate(headers, start=1):
        width = len(header) + 4
        for row in rows[:500]:
            value = row[idx - 1]
            if value is not None:
                width = max(width, min(len(str(value)) + 2, 70))
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_workbook(docs: list[AppDocumentation], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    used_sheets: set[str] = set()
    used_tables: set[str] = set()

    summary = wb.active
    summary.title = "Summary"
    used_sheets.add("summary")
    summary.append(["Dataverse Model-Driven App Documentation"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    for doc in docs:
        summary.append(["App", doc.app_name])
        summary.append(["Unique Name", doc.app_unique_name])
        summary.append(["App Id", doc.app_id])
        summary.append(["Environment", doc.environment_url])
        summary.append(["Generated (UTC)", doc.generated_at])
        summary.append(["Entities discovered", len(doc.entities)])
        summary.append(["Components documented", len(doc.components)])
        summary.append([])
        summary.append(["Component Type", "Count"])
        summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
        summary.cell(row=summary.max_row, column=2).font = Font(bold=True)
        for ctype, count in doc.counts_by_type().items():
            summary.append([ctype, count])
        if doc.warnings:
            summary.append([])
            summary.append(["Warnings"])
            summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
            for w in doc.warnings:
                summary.append([w])
        summary.append([])
        summary.append([])
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 80

    for doc in docs:
        ws = wb.create_sheet(sheet_name(doc.app_name, "Components", used_sheets))
        comp_rows = [
            [c.component_name, c.component_type, c.logical_name, c.parent_entity, c.component_id, c.discovery_source, c.details]
            for c in sorted(doc.components, key=lambda c: (c.component_type, c.parent_entity, c.component_name.lower()))
        ]
        _write_table(ws, COMPONENT_HEADERS, comp_rows, _table_name(f"{doc.app_unique_name or doc.app_name}_Components", used_tables))

        ws2 = wb.create_sheet(sheet_name(doc.app_name, "Entities", used_sheets))
        ent_rows = [
            [
                e.logical_name,
                e.display_name,
                e.discovery_source,
                "Yes" if e.is_custom_entity else "No",
                "; ".join(e.discovery_sources),
                "Yes" if e.is_activity else "No",
                e.schema_name,
                e.object_type_code,
                e.depth,
                "; ".join(e.referenced_by),
            ]
            for e in sorted(doc.entities.values(), key=lambda e: (e.depth, e.logical_name))
        ]
        _write_table(ws2, ENTITY_HEADERS, ent_rows, _table_name(f"{doc.app_unique_name or doc.app_name}_Entities", used_tables))

    wb.save(path)
    return path
