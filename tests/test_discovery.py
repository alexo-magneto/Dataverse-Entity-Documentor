"""End-to-end test of the walker against an in-memory fake Dataverse."""

import sys
import tempfile
import unittest
from pathlib import Path
from urllib.parse import parse_qs, urlparse

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from dataverse_app_documentor.config import DiscoveryConfig  # noqa: E402
from dataverse_app_documentor.discovery import AppDocumenter  # noqa: E402
from dataverse_app_documentor.metadata import MetadataCache  # noqa: E402
from tests.test_parsers import CHART_DATA, DASHBOARD, FORM, SITEMAP  # noqa: E402

APP_ID = "aaaaaaaa-0000-0000-0000-000000000001"
APP_UNIQUE = "aaaaaaaa-0000-0000-0000-0000000000ff"
SITEMAP_ID = "bbbbbbbb-0000-0000-0000-000000000001"
ACCOUNT_MD = "cccccccc-0000-0000-0000-000000000001"
ACCOUNT_FORM = "dddddddd-0000-0000-0000-000000000001"
CONTACT_QV = "a1b2c3d4-0000-0000-0000-000000000001"
DASH_ID = "2e3d0841-fa6d-4a6a-8b76-6a0b1b2c3d4e"
VIEW_ID = "00000000-0000-0000-00aa-000010001004"
CHART_ID = "c1"
BPF_ID = "eeeeeeee-0000-0000-0000-000000000001"


def label(text):
    return {"UserLocalizedLabel": {"Label": text}, "LocalizedLabels": [{"Label": text}]}


def ent(name, display, md, otc, custom=False, activity=False, notes=True, acts=True):
    return {
        "LogicalName": name, "DisplayName": label(display), "SchemaName": name.capitalize(), "EntitySetName": name + "s",
        "MetadataId": md, "ObjectTypeCode": otc, "IsCustomEntity": custom, "IsActivity": activity,
        "HasActivities": acts, "HasNotes": notes, "PrimaryIdAttribute": name + "id", "PrimaryNameAttribute": "name", "IsManaged": not custom,
    }


ENTITIES = [
    ent("account", "Account", ACCOUNT_MD, 1),
    ent("contact", "Contact", "cccccccc-0000-0000-0000-000000000002", 2),
    ent("opportunity", "Opportunity", "cccccccc-0000-0000-0000-000000000003", 3),
    ent("lead", "Lead", "cccccccc-0000-0000-0000-000000000004", 4),
    ent("systemuser", "User", "cccccccc-0000-0000-0000-000000000008", 8, notes=False, acts=False),
    ent("team", "Team", "cccccccc-0000-0000-0000-000000000009", 9, notes=False, acts=False),
    ent("email", "Email", "cccccccc-0000-0000-0000-000000004202", 4202, activity=True),
    ent("phonecall", "Phone Call", "cccccccc-0000-0000-0000-000000004210", 4210, activity=True),
    ent("task", "Task", "cccccccc-0000-0000-0000-000000004212", 4212, activity=True),
    ent("new_customactivity", "Custom Activity", "cccccccc-0000-0000-0000-000000010001", 10001, custom=True, activity=True),
    ent("activitypointer", "Activity", "cccccccc-0000-0000-0000-000000004200", 4200),
    ent("annotation", "Note", "cccccccc-0000-0000-0000-000000000005", 5),
    ent("opportunitysalesprocess", "Opportunity Sales Process", "cccccccc-0000-0000-0000-000000010002", 10002, custom=True),
]

ATTRS = {
    "account": [
        {"LogicalName": "name", "DisplayName": label("Account Name"), "SchemaName": "Name", "AttributeType": "String", "IsCustomAttribute": False},
        {"LogicalName": "primarycontactid", "DisplayName": label("Primary Contact"), "SchemaName": "PrimaryContactId", "AttributeType": "Lookup", "IsCustomAttribute": False},
        {"LogicalName": "ownerid", "DisplayName": label("Owner"), "SchemaName": "OwnerId", "AttributeType": "Owner", "IsCustomAttribute": False},
    ],
}
LOOKUPS = {
    "account": [
        {"LogicalName": "primarycontactid", "Targets": ["contact"]},
        {"LogicalName": "ownerid", "Targets": ["systemuser", "team"]},
    ],
}
RELS = {
    "account": {
        "OneToManyRelationships": [
            {"SchemaName": "contact_customer_accounts", "ReferencedEntity": "account", "ReferencingEntity": "contact", "ReferencingAttribute": "parentcustomerid"},
            {"SchemaName": "opportunity_customer_accounts", "ReferencedEntity": "account", "ReferencingEntity": "opportunity", "ReferencingAttribute": "customerid"},
        ],
        "ManyToOneRelationships": [],
        "ManyToManyRelationships": [],
    }
}


class FakeClient:
    """Answers the handful of Web API shapes the documenter uses."""

    environment_url = "https://fake.crm.dynamics.com"

    def __init__(self):
        self.request_count = 0
        self.calls = []

    def _filter(self, params):
        return (params or {}).get("$filter", "")

    def get_all(self, path, params=None):
        self.request_count += 1
        self.calls.append((path, params))
        flt = self._filter(params)
        if path == "EntityDefinitions":
            return ENTITIES
        if path.startswith("EntityDefinitions(") and path.endswith("/Attributes"):
            name = path.split("'")[1]
            return ATTRS.get(name, [])
        if path.endswith("LookupAttributeMetadata"):
            name = path.split("'")[1]
            return LOOKUPS.get(name, [])
        if path == "appmodules":
            return [{"appmoduleid": APP_ID, "appmoduleidunique": APP_UNIQUE, "name": "Sales Hub", "uniquename": "msdynce_saleshub", "description": "Test app"}] if ("Sales Hub" in flt or not flt) else []
        if path == "appmodulecomponents":
            if APP_UNIQUE not in flt:
                return []
            return [
                {"objectid": ACCOUNT_MD, "componenttype": 1, "rootcomponentbehavior": 0},
                {"objectid": SITEMAP_ID, "componenttype": 62},
                {"objectid": DASH_ID, "componenttype": 60},
                {"objectid": BPF_ID, "componenttype": 29},
            ]
        if path == "systemforms":
            forms = {
                ACCOUNT_FORM: {"formid": ACCOUNT_FORM, "name": "Account Main", "type": 2, "objecttypecode": "account", "formactivationstate": 1, "isdefault": True, "formxml": FORM, "uniquename": ""},
                DASH_ID: {"formid": DASH_ID, "name": "Sales Dashboard", "type": 0, "objecttypecode": "none", "formactivationstate": 1, "isdefault": False, "formxml": DASHBOARD},
                CONTACT_QV: {"formid": CONTACT_QV, "name": "Contact Quick View", "type": 6, "objecttypecode": "contact", "formactivationstate": 1, "isdefault": False, "formxml": "<form/>"},
            }
            if "objecttypecode eq 'account'" in flt:
                return [forms[ACCOUNT_FORM]]
            if "objecttypecode eq" in flt:
                return []
            return [f for fid, f in forms.items() if fid in flt]
        if path == "savedqueries":
            views = {VIEW_ID: {"savedqueryid": VIEW_ID, "name": "Active Contacts Subgrid View", "returnedtypecode": "contact", "querytype": 2, "fetchxml": '<fetch><entity name="contact"><link-entity name="systemuser" from="systemuserid" to="ownerid"/></entity></fetch>', "isdefault": False, "statecode": 0}}
            if "returnedtypecode eq 'account'" in flt:
                return [{"savedqueryid": "v-acc-1", "name": "Active Accounts", "returnedtypecode": "account", "querytype": 0, "fetchxml": '<fetch><entity name="account"/></fetch>', "isdefault": True, "statecode": 0}]
            if "returnedtypecode eq" in flt:
                return []
            return [v for vid, v in views.items() if vid in flt]
        if path == "savedqueryvisualizations":
            if "primaryentitytypecode eq 'account'" in flt:
                return [{"savedqueryvisualizationid": "ch-acc-1", "name": "Accounts by Owner", "primaryentitytypecode": "account", "datadescription": '<datadefinition><fetchcollection><fetch><entity name="account"><link-entity name="systemuser"/></entity></fetch></fetchcollection></datadefinition>'}]
            if "primaryentitytypecode eq" in flt:
                return []
            if "c1" in flt:
                return [{"savedqueryvisualizationid": "c1", "name": "Opportunities by Owner", "primaryentitytypecode": "opportunity", "datadescription": CHART_DATA}]
            return []
        if path == "webresources":
            return [{"webresourceid": "wr-1", "name": "new_/pages/help.html", "displayname": "Help Page", "webresourcetype": 1}]
        return []

    def get(self, path, params=None):
        self.request_count += 1
        self.calls.append((path, params))
        if path.startswith("EntityDefinitions("):
            name = path.split("'")[1]
            return {"LogicalName": name, **RELS.get(name, {})}
        raise AssertionError(f"unexpected get {path}")

    def get_by_id(self, entity_set, record_id, select):
        self.request_count += 1
        if entity_set == "sitemaps":
            return {"sitemapid": SITEMAP_ID, "sitemapname": "Sales Hub SiteMap", "sitemapnameunique": "msdynce_saleshub_sitemap", "sitemapxml": SITEMAP}
        if entity_set == "workflows":
            return {"workflowid": BPF_ID, "name": "Opportunity Sales Process", "uniquename": "opportunitysalesprocess", "primaryentity": "opportunity", "category": 4, "type": 1, "statecode": 1,
                    "clientdata": '{"steps":[{"entityName":"opportunity"},{"entityName":"lead"}]}'}
        return None

    def get_many_by_ids(self, entity_set, id_attribute, ids, select, chunk=20):
        ids = [i for i in dict.fromkeys(ids) if i]
        if not ids:
            return []
        flt = " or ".join(f"{id_attribute} eq {i}" for i in ids)
        return self.get_all(entity_set, {"$select": select, "$filter": flt})


class DiscoveryTests(unittest.TestCase):
    def _run(self, **overrides):
        cfg = DiscoveryConfig(**overrides)
        client = FakeClient()
        doc = AppDocumenter(client, MetadataCache(client), cfg).document_app("Sales Hub")
        return doc, client

    def test_full_walk(self):
        doc, _ = self._run()
        ents = doc.entities
        # Root + sitemap
        self.assertIn("AppModuleComponent", ents["account"].discovery_sources)
        self.assertIn("Sitemap", ents["account"].discovery_sources)
        self.assertEqual(ents["contact"].discovery_sources[0], "Sitemap")
        # Discovered from the account form
        self.assertIn("Lookup Field", ents["contact"].discovery_sources)
        self.assertIn("Form Subgrid", ents["contact"].discovery_sources)
        self.assertIn("Quick View Form", ents["contact"].discovery_sources)
        self.assertIn("Lookup Field", ents["systemuser"].discovery_sources)
        self.assertIn("Lookup Field", ents["team"].discovery_sources)
        self.assertIn("Form Navigation (Relationship)", ents["opportunity"].discovery_sources)
        # Timeline -> activity entities + notes
        for act in ("email", "phonecall", "task", "new_customactivity", "activitypointer", "annotation"):
            self.assertIn("Timeline", ents[act].discovery_sources, act)
        self.assertTrue(ents["new_customactivity"].is_custom_entity)
        self.assertTrue(ents["email"].is_activity)
        # Dashboard
        self.assertIn("Dashboard", ents["opportunity"].discovery_sources)
        self.assertIn("Dashboard", ents["lead"].discovery_sources)
        # BPF
        self.assertIn("Business Process Flow", ents["opportunitysalesprocess"].discovery_sources)
        self.assertIn("Business Process Flow", ents["lead"].discovery_sources)
        # Views/charts discovered through subgrid + chart ids
        self.assertIn("View Link-Entity", ents["systemuser"].discovery_sources)
        self.assertIn("Chart", ents["systemuser"].discovery_sources)

        types = doc.counts_by_type()
        self.assertEqual(types["Model-Driven App"], 1)
        self.assertEqual(types["Site Map"], 1)
        self.assertEqual(types["Sitemap SubArea"], 4)
        self.assertEqual(types["Entity"], 1)
        self.assertEqual(types["Main Form"], 1)
        self.assertEqual(types["Quick View Form"], 1)
        self.assertEqual(types["Dashboard"], 1)
        self.assertEqual(types["Form Tab"], 1)
        self.assertEqual(types["Attribute (Form Field)"], 3)
        self.assertEqual(types["Form Subgrid"], 1)
        self.assertEqual(types["Form Timeline"], 1)
        self.assertEqual(types["Business Process Flow"], 1)
        self.assertEqual(types["Web Resource"], 1)
        self.assertEqual(types["View"], 2)   # account view + subgrid view
        self.assertEqual(types["Chart"], 2)  # account chart + dashboard chart
        # depth-1 entities were not walked (traversal depth 1)
        self.assertEqual(ents["contact"].depth, 0)  # sitemap root
        self.assertEqual(ents["team"].depth, 1)
        self.assertFalse(doc.warnings, doc.warnings)

    def test_depth_zero_skips_forms(self):
        doc, _ = self._run(traversal_depth=0)
        self.assertNotIn("Main Form", doc.counts_by_type())
        self.assertNotIn("team", doc.entities)
        self.assertIn("Dashboard", doc.counts_by_type())  # dashboards are app components, always documented

    def test_ignore_lookup_and_exclude(self):
        doc, _ = self._run(ignore_lookup_attributes=["ownerid"], exclude_entities=["annotation"])
        self.assertNotIn("team", doc.entities)
        self.assertNotIn("annotation", doc.entities)
        self.assertIn("contact", doc.entities)

    def test_excel_output(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")
        from dataverse_app_documentor.excel import write_workbook
        doc, _ = self._run()
        with tempfile.TemporaryDirectory() as tmp:
            path = write_workbook([doc], Path(tmp) / "out.xlsx")
            wb = openpyxl.load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Summary", "Sales Hub - Components", "Sales Hub - Entities"])
            ws = wb["Sales Hub - Entities"]
            headers = [c.value for c in ws[1]]
            self.assertEqual(headers[:4], ["Logical Name", "Display Name", "Discovery Source", "IsCustomEntity"])
            self.assertEqual(ws.max_row - 1, len(doc.entities))
            wsc = wb["Sales Hub - Components"]
            self.assertEqual([c.value for c in wsc[1]][:2], ["Component Name", "Component Type"])
            self.assertEqual(wsc.max_row - 1, len(doc.components))


if __name__ == "__main__":
    unittest.main()
